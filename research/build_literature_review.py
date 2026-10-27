"""Build research/literature_review.xlsx — one row per paper or dataset.

Column layout follows the department's expected format (ID, category, paper,
authors, year, venue, links, datasets, main work, method, headline results as
reported, hardware tested, limitation, how our project addresses it,
verification).  Run:  python3 research/build_literature_review.py
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "literature_review.xlsx")

TITLE = ("Literature review: lightweight IoT intrusion detection on CICIoT2023, "
         "cross-dataset generalisation and edge deployment")
SUBTITLE = ("Each row was checked against the paper, abstract, dataset page or repository page "
            "named in the Verification column. Numbers are as reported by the authors; "
            "\"not copied here\" means we did not transcribe them.")

HEADERS = [
    "ID", "Category", "Paper title", "Authors", "Year", "Venue", "Paper link",
    "Dataset(s) used", "Dataset link(s)", "Main work", "Model / method",
    "Headline results (as reported)", "Hardware / deployment tested",
    "Bottleneck / limitation", "How our project addresses or enhances it", "Verification",
]
WIDTHS = [6, 20, 34, 28, 6, 24, 30, 28, 34, 46, 30, 44, 30, 44, 46, 18]

L = {
    "ciciot": "https://www.unb.ca/cic/datasets/iotdataset-2023.html",
    "iomt": "https://www.unb.ca/cic/datasets/iomt-dataset-2024.html",
    "ids2017": "https://www.unb.ca/cic/datasets/ids-2017.html",
    "ids2018": "https://www.unb.ca/cic/datasets/ids-2018.html",
    "nsl": "https://www.unb.ca/cic/datasets/nsl.html",
    "nbaiot": "https://archive.ics.uci.edu/dataset/442/detection+of+iot+botnet+attacks+n+baiot",
    "edge": "https://ieee-dataport.org/documents/edge-iiotset-new-comprehensive-realistic-cyber-security-dataset-iot-and-iiot-applications",
    "ton": "https://research.unsw.edu.au/projects/toniot-datasets",
    "bot": "https://research.unsw.edu.au/projects/bot-iot-dataset",
    "unsw": "https://research.unsw.edu.au/projects/unsw-nb15-dataset",
    "tabular": "https://www.unb.ca/cic/datasets/tabular-iot-attack-2024.html",
    "nf": "https://www.cyber.uq.edu.au/node/824",
    "rtiot": "https://archive.ics.uci.edu/dataset/942/rt-iot2022",
}

ROWS = [
    # ---------------------------------------------------------------- datasets
    ("P01", "Dataset paper (CICIoT2023)",
     "CICIoT2023: A Real-Time Dataset and Benchmark for Large-Scale Attacks in IoT Environment",
     "E. C. P. Neto, S. Dadkhah, R. Ferreira, A. Zohourian, R. Lu, A. A. Ghorbani", 2023,
     "Sensors 23(13):5941", "https://www.mdpi.com/1424-8220/23/13/5941",
     "CICIoT2023 (created in this paper): 105 real IoT devices, 33 attacks in 7 categories, 46 flow features, 46.7M records",
     L["ciciot"],
     "Builds a real IoT testbed, runs 33 attacks, extracts 46 statistical features with CIC's own tool and publishes "
     "baselines (Logistic Regression, Perceptron, AdaBoost, Random Forest, DNN) for binary, 8-class and 34-class tasks.",
     "Baseline ML/DL classifiers on hand-crafted flow features",
     "Random Forest is the strongest baseline in all three settings; rare classes (Web-based, BruteForce) are detected "
     "much worse than DDoS/DoS. Exact per-setting figures not copied here.",
     "None on device; dataset captured on a 105-device testbed",
     "72.8% of records are DDoS, so accuracy on the natural mix mostly reflects DDoS. Features come from CIC's own "
     "extractor (not CICFlowMeter), which complicates mixing with other datasets. No cross-dataset test.",
     "Our home dataset (canonical 46.7M-row release). We report both a natural-mix and a class-balanced test set with "
     "macro-F1 and MCC, so published numbers can be reproduced and stress-tested.",
     "Abstract and dataset page read; data loaded (46,686,579 rows checked)"),

    ("P02", "Dataset paper (CICIoMT2024)",
     "CICIoMT2024: A benchmark dataset for multi-protocol security assessment in IoMT",
     "S. Dadkhah, E. C. P. Neto, R. Ferreira, R. C. Molokwu, S. Sadeghi, A. A. Ghorbani", 2024,
     "Internet of Things (Elsevier)",
     "https://www.researchgate.net/publication/383572709_CICIoMT2024_A_benchmark_dataset_for_multi-protocol_security_assessment_in_IoMT",
     "CICIoMT2024 (created in this paper): 40 medical-IoT devices (25 real, 15 simulated), 18 attacks over Wi-Fi, MQTT and Bluetooth",
     L["iomt"],
     "Medical-IoT testbed with the same CIC feature extractor as CICIoT2023 (45 features); attacks include DDoS, DoS, "
     "Recon, MQTT floods and ARP spoofing; baseline classifiers for binary, 6-class and 19-class tasks.",
     "Baseline ML/DL classifiers",
     "Baselines reported for binary, 6-class and 19-class tasks; figures not copied here (abstract only).",
     "None",
     "The attack CSVs are 97.4% attacks, so plain accuracy is uninformative. Evaluated in isolation in the paper.",
     "Our unseen target network for E2/E3. 45 of the 46 CICIoT2023 features align by name, so the exact trained models "
     "run unchanged. 8.7M flows loaded.",
     "Abstract and dataset page read; data loaded and checked"),

    # ------------------------------------------------- lightweight CICIoT2023 models
    ("P03", "Lightweight IDS on CICIoT2023",
     "Lightweight MLP-Based Feature Extraction with Linear Classifier for Intrusion Detection System in Internet of Things",
     "J. Chandroth, J. Ali", 2026, "Electronics 15(8):1604", "https://doi.org/10.3390/electronics15081604",
     "CICIDS2017, NSL-KDD, CICIoT2023",
     f"{L['ids2017']}\n{L['nsl']}\n{L['ciciot']}",
     "A two-layer MLP embedding backbone followed by a linear softmax head, designed to minimise parameters and FLOPs "
     "while keeping multi-class accuracy.",
     "2-layer MLP + linear classifier",
     "Accuracy 99.85% (CICIDS2017), 99.21% (NSL-KDD), 98.45% (CICIoT2023); about 14,662 parameters and 28,928 FLOPs "
     "on CICIoT2023.",
     "Not stated in the abstract",
     "Each dataset is tested only against itself; accuracy is on the natural class mix; no code released; layer widths "
     "are not in the abstract.",
     "Rebuilt as MLP14k (14,336 params). We reproduce 99.3% binary accuracy on the natural mix, but get 76% on the "
     "balanced 8-class test and MCC 0.37 on CICIoMT2024.",
     "Abstract read (full text blocked at MDPI)"),

    ("P04", "Lightweight IDS on CICIoT2023",
     "Lightweight intrusion detection model based on CNN and knowledge distillation",
     "L.-H. Wang, Q. Dai, T. Du, L.-F. Chen", 2024, "Applied Soft Computing (doi 10.1016/j.asoc.2024.112118)",
     "https://doi.org/10.1016/j.asoc.2024.112118",
     "CICIoT2023 (other datasets not captured)", L["ciciot"],
     "A CNN teacher is distilled into a very small student CNN; the paper compares activation functions for the student "
     "and reports the accuracy/efficiency trade-off.",
     "CNN teacher -> tiny CNN student via knowledge distillation",
     "Student with 1,920 parameters reaches about 96% accuracy; the ReLU student gives the best balance between "
     "performance and efficiency.",
     "Not captured",
     "Single-dataset evaluation; no transfer test; no code.",
     "Rebuilt a 1,940-param student trained with KD (T=4, alpha=0.7) from a 14.9k-param teacher. Lowest in-distribution "
     "scores of our set, but the only model whose cross-dataset transfer does not degrade with longer training.",
     "Abstract read"),

    ("P05", "Lightweight IDS on CICIoT2023",
     "A lightweight intrusion detection method for IoT based on deep learning and dynamic quantization",
     "Z. Wang, H. Chen, S. Yang, X. Luo, D. Li, J. Wang", 2023, "PeerJ Computer Science 9:e1569",
     "https://doi.org/10.7717/peerj-cs.1569",
     "CICIDS2017, N-BaIoT, CICIoT2023 (Recon and Mirai classes only)",
     f"{L['ids2017']}\n{L['nbaiot']}\n{L['ciciot']}",
     "A dual-hidden-layer DNN feeds a bidirectional LSTM (DL-BiLSTM); incremental PCA for dimensionality reduction, "
     "Optuna tuning, then post-training dynamic INT8 quantisation of the LSTM.",
     "DNN + BiLSTM, dynamic INT8 quantisation",
     "Accuracy 99.67% (CICIDS2017), 99.98% (N-BaIoT), 93.13% (CICIoT2023 subset); 1,988 parameters and 36.5 KB on "
     "CICIoT2023.",
     "Intel Core i5-7200U laptop, 12 GB RAM; no edge device",
     "CICIoT2023 is evaluated only on Recon and Mirai; layer widths are not given; the authors say dynamic quantisation "
     "did not work well on Linear layers; no cross-dataset test.",
     "Rebuilt as DLBiLSTM (11,432 params because widths are not given) and evaluated on all 8 classes. It is our best "
     "cross-dataset model at 4 epochs (MCC 0.830 +/- 0.013 on CICIoMT2024), and we confirm INT8 costs 6-10 macro-F1 points.",
     "Full text read"),

    ("P06", "Lightweight IDS / edge (code artifact)",
     "reliable-edge-iot-ids: reliability-aware lightweight network-flow intrusion detection for IoT/IIoT edge networks",
     "F. Keskin", 2026, "GitHub research artifact (manuscript in submission)",
     "https://github.com/fesihkeskin/reliable-edge-iot-ids",
     "CICIoT2023 (primary), Edge-IIoTset (validation)", f"{L['ciciot']}\n{L['edge']}",
     "Companion code for an unpublished study: Logistic Regression, Random Forest, LightGBM and a TinyMLP with pruning "
     "and dynamic INT8; binary/multiclass, leave-one-attack-out zero-day test, calibration, compression, throughput.",
     "LR / RF / LightGBM / TinyMLP + pruning + INT8",
     "None on the repository page (the results folder is not included).",
     "Not reported; notes that latency depends on hardware",
     "Unpublished; no numbers; raw data not redistributed; 0 stars.",
     "Adopted the TinyMLP + 50% pruning + INT8 recipe in our own implementation (7,688 params). Pruning cost 0.013 "
     "macro-F1; INT8 cost 0.057.",
     "Repository page read"),

    ("P07", "Edge IDS on CICIoT2023",
     "A Lightweight Two-Stage Intrusion Detection Framework Optimized for Edge-Based IoT Environments",
     "C.-W. Kuo, C.-X. Wu", 2026, "Computers, Materials & Continua", "https://doi.org/10.32604/cmc.2026.076767",
     "CICIoT2023", L["ciciot"],
     "A two-stage lightweight detection framework for edge IoT that cuts training and testing time compared with "
     "single-stage approaches.",
     "Two-stage lightweight classifier",
     "99.67% accuracy on the full CICIoT2023; about 64.5% less training/testing time than single-stage; validated on a "
     "Raspberry Pi.",
     "Raspberry Pi (as reported)",
     "Single dataset; accuracy on the natural class mix; no cross-dataset test.",
     "Our tree baselines also reach 99.3-99.5% on the natural mix yet collapse to MCC 0.07-0.17 on CICIoMT2024, which "
     "shows what such headline numbers hide.",
     "Abstract read"),

    ("P08", "Edge IDS on CICIoT2023",
     "A Lightweight Intrusion Detection System Using Convolutional Neural Network and Long Short-Term Memory in Fog Computing",
     "H. Alzahrani, T. Sheltami, A. Barnawi, M. Imam, A. Yaser", 2024, "Computers, Materials & Continua 80(3)",
     "https://doi.org/10.32604/cmc.2024.054203",
     "CICIoT2023, KDD-99, NSL-KDD", f"{L['ciciot']}\n{L['nsl']}",
     "A hybrid CNN-LSTM anomaly IDS for fog nodes, evaluated on latency, energy, false-alarm rate and detection rate.",
     "CNN-LSTM",
     "Accuracy above 92%, false-alarm rate below 0.38%, maximum power 6.12 W on the Raspberry Pi.",
     "Raspberry Pi as fog node",
     "Two of the three datasets are over 20 years old; no cross-dataset test.",
     "Template for our edge-cost protocol (E5); the Raspberry Pi measurement is our planned next step.",
     "Abstract read"),

    ("P09", "Edge IDS on CICIoT2023",
     "A resource efficient IoT intrusion detection model using hybrid feature selection for edge computing",
     "M. Z. Khan, M. Alam, I. Alam, M. U. Bokhari, Z. Ashraf, M. Z. Khan, F. S. Masoodi", 2026,
     "Scientific Reports 16:24187", "https://doi.org/10.1038/s41598-026-59156-9",
     "Own Raspberry Pi capture, TON_IoT, CICIDS2018, CICIoT2023",
     f"{L['ton']}\n{L['ids2018']}\n{L['ciciot']}",
     "BGL-RID: Boruta then greedy forward feature selection (Extra Trees as evaluator), SMOTE, and a LightGBM "
     "classifier; proposes an 'efficiency ratio' that combines accuracy, FNR, time and memory.",
     "Feature selection + LightGBM",
     "CICIoT2023 accuracy 98.88%, F1 98.97%, 3.62 s, 17.6 MB; TON_IoT binary 99.72%; CICIDS2018 99.94%.",
     "Intel Core i5 laptop; the Raspberry Pi 4 only captured traffic",
     "The Pi data has proxy labels from Isolation Forest, not real attack labels; no cross-dataset test; the model is a "
     "tree ensemble.",
     "Our LightGBM baseline matches its in-distribution accuracy (99.5%) but drops to MCC 0.17 on CICIoMT2024 and "
     "loses the most under feature noise.",
     "Full text read"),

    ("P10", "Federated / 'TinyML' IDS",
     "Lightweight and Energy-Aware Intrusion Detection for Industrial IoT Using TinyML and Edge AI",
     "L. Nassef, M. I. Alghamdi, S. Ben Chaabane, Q. Abbas, W. M. Alawad, O. H. Albalawi, O. I. Alqaisi, B. Fakieh", 2026,
     "Scientific Reports 16:18524", "https://doi.org/10.1038/s41598-026-50690-0",
     "Edge-IIoTset, CICIoT2023, RT-IoT2022", f"{L['edge']}\n{L['ciciot']}\n{L['rtiot']}",
     "Graph attention network + BiGRU with Grey Wolf hyperparameter tuning, trained with FedAvg over 10 simulated edge "
     "nodes; attention weights used for explainability.",
     "GAT + BiGRU, federated",
     "5-fold accuracy / F1: Edge-IIoTset 0.950 / 0.940, CICIoT2023 0.930 / 0.920, RT-IoT2022 0.940 / 0.930; 8-12 MB of "
     "weights; 120-180 ms per sample on a Raspberry Pi 4 (preliminary).",
     "Workstation (Xeon + RTX 3080); preliminary Raspberry Pi 4 timing; no live deployment; TinyML not actually implemented",
     "Heavy graph construction; 300-500 MB RAM at inference; adversarial test is only 10% Gaussian noise; energy not measured.",
     "Our models are 11-60 KB and 0.013-0.17 ms per flow, an honest lightweight baseline; our drift tests cover noise, "
     "feature dropout and scale shift.",
     "Full text read"),

    # ------------------------------------------------------ generalisation studies
    ("P11", "Cross-dataset generalisation",
     "On the Cross-Dataset Generalization of Machine Learning for Network Intrusion Detection",
     "M. Cantone, C. Marrocco, A. Bria", 2024, "arXiv:2402.10974", "https://arxiv.org/abs/2402.10974",
     "Four NIDS datasets (see paper)", "See paper",
     "Trains ML-NIDS classifiers on one dataset and tests them on the others, measuring within- versus cross-dataset MCC.",
     "Several ML classifiers",
     "Average MCC about 94% within-dataset versus about 29% cross-dataset.",
     "None",
     "Generic classifiers rather than the specific published lightweight IoT models; no edge-cost dimension.",
     "We run the same test on four named published lightweight models plus two tree baselines, and add drift, "
     "synthetic-data and edge-cost arms.",
     "Abstract read"),

    ("P12", "Cross-dataset generalisation",
     "Assessing Generalisation Capability of Machine Learning Models for Intrusion Detection",
     "M. Z. Hossain, M. A. R. Khan, M. R. Islam, S. M. S. Islam, T. Gedeon", 2026, "arXiv:2605.04407",
     "https://arxiv.org/abs/2605.04407",
     "UNSW-NB15, TON_IoT", f"{L['unsw']}\n{L['ton']}",
     "Random Forest, Logistic Regression and Naive Bayes trained on one dataset and tested on the other.",
     "RF / LR / Naive Bayes",
     "RF 95.08% (UNSW-NB15) and 99.79% (TON_IoT) within-dataset; under 40% accuracy cross-dataset.",
     "None",
     "Classical models only; binary task; accuracy as the metric.",
     "We add four neural architectures and show that the ranking inverts under domain shift: trees collapse, the "
     "neural models survive.",
     "Abstract read"),

    ("P13", "Cross-dataset generalisation",
     "Distributional Drift in IoT Intrusion Detection Systems: Implications for Cross-Dataset Generalisation",
     "K. K. Eren, K. Kucuk, R. A. A. Saleh, M. Z. Konyar, O. M. Hardy, S. A. Khan", 2026, "Electronics 15(11):2307",
     "https://doi.org/10.3390/electronics15112307",
     "BoT-IoT, ToN-IoT, UNSW-NB15", f"{L['bot']}\n{L['ton']}\n{L['unsw']}",
     "Attributes cross-dataset failure of IoT IDS models to distributional drift and dataset-specific bias between "
     "the three datasets.",
     "ML classifiers + drift analysis",
     "Sharp cross-dataset drops reported; exact figures not copied here.",
     "None",
     "Diagnoses the problem without testing fixes.",
     "E4 quantifies each model's drift sensitivity (trees worst); E3 tests the fix, combined training, and finds it "
     "almost free.",
     "Abstract read"),

    ("P14", "Cross-dataset generalisation",
     "Dataset-centric evaluation of federated intrusion detection models in IoT networks",
     "M. A. Bilal, I. U. Islam, S. Idrees, M. Qasim, M. J. Khan, J. Khan", 2026, "Scientific Reports 16:2683",
     "https://doi.org/10.1038/s41598-025-32567-w",
     "Edge-IIoTset, CIC-IoT2023, TII-SSRC-23", f"{L['edge']}\n{L['ciciot']}",
     "Federated LSTM and Transformer models with FedAvg / FedProx / FedNova, evaluated in-domain, cross-dataset and as "
     "a combined federation on 40 common features with family-level labels.",
     "Federated LSTM / Transformer",
     "In-domain macro-F1 up to 96.5%; up to a 30-point macro-F1 loss on unseen environments (abstract); the joint "
     "federation reaches about 88-90% F1 on every test set.",
     "None (simulated clients)",
     "No domain adaptation; the authors note accuracy is inflated by the benign class; with sequence length 1 the LSTM "
     "is effectively feed-forward.",
     "E3 confirms the multi-dataset effect at model level: training on a 50/50 mix lifts CICIoMT2024 F1 to 0.96-0.97 at "
     "a 0.02-0.03 cost on the home dataset.",
     "Full text read"),

    ("P15", "Cross-dataset generalisation",
     "A Cross-Dataset based Zero-Day Intrusion Detection System by Integrating Siamese Network and Reinforcement Learning",
     "M. M. Hossain, S. D. Turja, S. Tasnim, M. F.-U.-A. Juboraj, M. I. Hossain", 2026, "arXiv:2609.26115",
     "https://arxiv.org/abs/2609.26115",
     "CICIoT2023, CIC-BCCC-NRC TabularIoTAttacks-2024", f"{L['ciciot']}\n{L['tabular']}",
     "Siamese-network anomaly correlation combined with a PPO reinforcement-learning agent for adaptive defence; trained "
     "on one dataset and tested on the other.",
     "Siamese network + PPO agent",
     "99.07% accuracy on unknown attacks; 93.94% zero-day detection ratio.",
     "None",
     "Complex two-component system; accuracy-centric metrics on attack-heavy data.",
     "Precedent for our planned CICIoT2023 -> TabularIoTAttacks-2024 extension; we report MCC and false-positive rate "
     "instead of accuracy.",
     "Abstract read"),

    ("P16", "Cross-dataset generalisation (lightweight)",
     "Cross-Domain Generalization Failure in Lightweight Intrusion Detection Models for IIoT Networks",
     "M. A. Hakim, M. S. Uddin, T. I. Anis", 2026, "arXiv:2607.00553", "https://arxiv.org/abs/2607.00553",
     "Three IIoT datasets (not named in the abstract)", "See paper",
     "Four lightweight models trained on one IIoT dataset and tested unchanged on two others using only shared features; "
     "studies shortcut features such as port categories.",
     "Four lightweight classifiers",
     "The most influential port-category feature appears 96-435 times more often in source-domain attack traffic than "
     "in the two target domains.",
     "None",
     "Datasets and models are not named in the abstract.",
     "CICIoT2023's 46 features contain no IPs or ports, so this shortcut is absent by design; the collapse we measure is "
     "therefore not a port artefact.",
     "Abstract read"),

    ("P17", "Standard feature sets",
     "Towards a Standard Feature Set for Network Intrusion Detection System Datasets",
     "M. Sarhan, S. Layeghy, M. Portmann", 2022, "Mobile Networks and Applications (arXiv:2101.11315)",
     "https://arxiv.org/abs/2101.11315",
     "NF-UNSW-NB15-v2, NF-BoT-IoT-v2, NF-ToN-IoT-v2, NF-CSE-CIC-IDS2018-v2, NF-UQ-NIDS-v2", L["nf"],
     "Proposes a standard 43-feature NetFlow set and regenerates five public datasets with it so models can be compared "
     "and transferred across datasets.",
     "Feature-set standardisation; Extra Trees evaluation",
     "Extra Trees on the 43-feature v2 datasets clearly beats the 8-feature v1 versions.",
     "None",
     "CICIoT2023 is not released in NetFlow format, so this set cannot be used directly for it.",
     "Motivates our unified-feature-space route; TabularIoTAttacks-2024 plays the same role for CIC-style features.",
     "Abstract read"),

    ("P18", "Dataset (aggregated)",
     "CIC-BCCC-NRC TabularIoTAttacks-2024",
     "Canadian Institute for Cybersecurity (UNB), BCCC (York University), National Research Council Canada", 2024,
     "Dataset release", L["tabular"],
     "Nine source datasets: IoT Network Intrusion 2019, MQTT-IoT-IDS-2020, TON_IoT, UQ IoT 2022, CIC IoT 2022, "
     "Edge-IIoTset, CICIoT2023, ACI-IoT-2023, CICIoMT2024", L["tabular"],
     "Re-extracts nine IoT datasets with CICFlowMeter into one tabular schema with normal/malicious labels; over one "
     "million records.",
     "Dataset aggregation",
     "Not applicable",
     "None",
     "Gated behind a registration form; CICFlowMeter features differ from CICIoT2023's native 46, so models must be "
     "retrained on the unified schema.",
     "Planned third dataset for E2/E3: train on its CICIoT2023 slice, test on the other eight.",
     "Dataset page read"),

    ("P19", "Dataset (IIoT)",
     "Edge-IIoTset: A New Comprehensive Realistic Cyber Security Dataset of IoT and IIoT Applications for Centralized "
     "and Federated Learning",
     "M. A. Ferrag, O. Friha, D. Hamouda, L. Maglaras, H. Janicke", 2022, "IEEE Access 10:40281-40306",
     "https://doi.org/10.1109/ACCESS.2022.3165809",
     "Edge-IIoTset (created in this paper): 61 features, 14 attacks in 5 categories", L["edge"],
     "Seven-layer IoT/IIoT testbed, 14 attacks, 61 features; baselines for centralised and federated learning.",
     "Baseline ML/DL, centralised and federated",
     "Centralised and federated baselines reported; figures not copied here.",
     "Testbed devices; no edge inference benchmark",
     "Different feature schema from CICIoT2023.",
     "Candidate third dataset; it is the second dataset in P06, P10 and P14.",
     "Abstract read"),

    # ---------------------------------------------------------- synthetic data
    ("P20", "Synthetic data for IDS",
     "Conditional Tabular Generative Adversarial Based Intrusion Detection System for Detecting DDoS and DoS Attacks on "
     "the Internet of Things Networks",
     "B. A. Alabsi, M. Anbar, S. D. A. Rihan", 2023, "Sensors 23(12):5644", "https://doi.org/10.3390/s23125644",
     "Bot-IoT", L["bot"],
     "CTGAN generates synthetic DDoS/DoS samples to balance Bot-IoT; several shallow and deep classifiers are then "
     "trained on the augmented data.",
     "CTGAN + ML/DL classifiers",
     "Augmentation improved accuracy, precision, recall and F1 of the classifiers; figures not copied here.",
     "None",
     "Single dataset; synthetic data used only for augmentation, not as a test condition.",
     "We use synthetic data as a test condition (train-real/test-synthetic and train-synthetic/test-real) with a "
     "Gaussian copula; CTGAN is the planned second generator.",
     "Abstract read"),

    # ------------------------------------------------- synopsis literature (kept)
    ("P21", "Survey (lightweight IoT IDS)",
     "A systematic literature review of recent lightweight detection approaches leveraging machine and deep learning "
     "mechanisms in Internet of Things networks",
     "G. AL Mukhaini, M. Anbar, S. Manickam, T. A. Al-Amiedy, A. Al Momani", 2024,
     "Journal of King Saud University - Computer and Information Sciences 36(1):101866",
     "https://doi.org/10.1016/j.jksuci.2023.101866",
     "Survey (no experiments)", "Not applicable",
     "Systematic review of lightweight ML/DL detection approaches for IoT networks.",
     "Review",
     "Finds that device resources, not accuracy, are the binding constraint, and that most papers never measure "
     "latency or memory on real hardware.",
     "None",
     "Review only; no experiments.",
     "E5 measures size, latency, INT8 and pruning for every model; the Raspberry Pi run is planned.",
     "Abstract read"),

    ("P22", "Lightweight IoT IDS (autoencoder)",
     "A Lightweight Intelligent Network Intrusion Detection System Using One-Class Autoencoder and Ensemble Learning for IoT",
     "W. Yao, L. Hu, Y. Hou, X. Li", 2023, "Sensors 23(8):4141", "https://doi.org/10.3390/s23084141",
     "Not captured", "Not captured",
     "A one-class bidirectional GRU autoencoder trained on benign traffic only, with ensemble learning to place unknown "
     "attacks near the closest known category.",
     "One-class BiGRU autoencoder + ensemble",
     "Not copied here.",
     "Not captured",
     "Needs clean benign data; edge cost not the focus.",
     "Origin of the synopsis's bidirectional, benign-only idea; our DL-BiLSTM reproduction keeps the bidirectional design.",
     "Abstract read"),

    ("P23", "Fog-layer IDS (BiLSTM)",
     "Optimized deep autoencoder and BiLSTM for intrusion detection in IoTs-Fog computing",
     "A. Alqahtani", 2025, "Multimedia Tools and Applications 84(8):4907-4943", "DOI not captured",
     "CICIDS2017, BoT-IoT", f"{L['ids2017']}\n{L['bot']}",
     "A deep autoencoder paired with a BiLSTM at the fog layer, tuned with the Sparrow Search Algorithm.",
     "Autoencoder + BiLSTM, metaheuristic tuning",
     "Higher detection rate and lower latency than the compared baselines.",
     "Fog layer (as reported)",
     "No cross-dataset test; metaheuristic tuning adds cost.",
     "Supports the BiLSTM-at-the-edge direction; we add cross-dataset and quantisation tests.",
     "Abstract read"),

    ("P24", "Lightweight IoT IDS (autoencoder)",
     "An automated intrusion detection system in IoT system using Attention based Deep Bidirectional Sparse Auto Encoder model",
     "K. Swathi, G. Hima Bindu", 2024, "Knowledge-Based Systems 305:112633",
     "https://doi.org/10.1016/j.knosys.2024.112633",
     "Not captured", "Not captured",
     "An attention layer before a bidirectional sparse autoencoder compresses features before classification; compared "
     "against unidirectional baselines.",
     "Attention + bidirectional sparse autoencoder",
     "Not copied here.",
     "Not captured",
     "Single-dataset; no transfer test.",
     "Its bidirectional-versus-unidirectional ablation is the kind of controlled comparison we run across all models.",
     "Abstract read"),
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Literature review"

    navy = "0B1F3A"
    thin = Side(style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = TITLE
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=navy)
    ws["A2"] = SUBTITLE
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="555555")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))

    hr = 4
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[hr].height = 30

    link_cols = {HEADERS.index("Paper link") + 1, HEADERS.index("Dataset link(s)") + 1}
    for r, row in enumerate(ROWS, hr + 1):
        fill = PatternFill("solid", fgColor="F2F4F7") if r % 2 else None
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(name="Calibri", size=9, bold=(c in (1, 3)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if fill:
                cell.fill = fill
            if c in link_cols and isinstance(val, str) and val.startswith("http"):
                first = val.split("\n")[0]
                cell.hyperlink = first
                cell.font = Font(name="Calibri", size=9, color="1F5FBF", underline="single")
        ws.row_dimensions[r].height = 150

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=4)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(HEADERS))}{hr + len(ROWS)}"

    wb.save(OUT)
    print(f"wrote {OUT} ({len(ROWS)} rows)")


if __name__ == "__main__":
    main()
